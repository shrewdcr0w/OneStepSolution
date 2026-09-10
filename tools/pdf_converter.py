import os
import time
import subprocess
import shutil
from PIL import Image
from PyPDF2 import PdfMerger
import pdfplumber
from fpdf import FPDF
from docx import Document

# Automatically checks common location of LibreOffice
def find_libreoffice():
    # 1. Check if it's already in the System PATH (Linux/Mac/Manual Windows setup)
    cmd_to_find = "soffice.exe" if os.name == 'nt' else "soffice"
    env_path = shutil.which(cmd_to_find) or shutil.which("libreoffice")
    if env_path:
            return env_path

    # 2. Check common Windows locations
    if os.name == 'nt':
        windows_paths = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"
        ]
        for p in windows_paths:
            if os.path.exists(p):
                return p
            
    # 3. Check Mac location
    mac_path = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if os.path.exists(mac_path):
        return mac_path
    
    # 4. LINUX SPECIFIC FALLBACKS (If not in PATH)
    linux_paths = [
        "/usr/bin/libreoffice",
        "/usr/bin/soffice",
        "/snap/bin/libreoffice"
    ]
    for p in linux_paths:
        if os.path.exists(p):
            return p

    return None

# Step A: Try automatic search
LIBRE_OFFICE_PATH = find_libreoffice()

# Step B: If automatic search failed (returned None), use your manual D: drive paths
if LIBRE_OFFICE_PATH is None:
    # Check your specific D: drive locations
    manual_path_1 = r'D:\LIBRA\program\soffice.exe'
    manual_path_2 = r'D:\LibreOffice\program\soffice.exe'
    
    if os.path.exists(manual_path_1):
        LIBRE_OFFICE_PATH = manual_path_1
    elif os.path.exists(manual_path_2):
        LIBRE_OFFICE_PATH = manual_path_2

# Step C: Define the helper function using the final result
def libreoffice_available():
    """Checks if the configured path actually exists."""
    return LIBRE_OFFICE_PATH is not None and os.path.exists(LIBRE_OFFICE_PATH)

# -----------------------------
# PDF → TXT
# -----------------------------
def pdf_to_txt(pdf_file, output_file):
    try:
        with pdfplumber.open(pdf_file) as pdf:
            text = "".join([page.extract_text() or "" for page in pdf.pages])
        with open(output_file, "w", encoding="utf-8") as f:
            f.write(text)
    except Exception as e:
        print(f"❌ PDF to TXT Error: {e}")
        raise e

# -----------------------------
# PDF → DOCX (Text-based)
# -----------------------------
def pdf_to_docx(pdf_file, output_file):
    try:
        doc = Document()
        with pdfplumber.open(pdf_file) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    for line in text.split("\n"):
                        doc.add_paragraph(line)
        doc.save(output_file)
    except Exception as e:
        print(f"❌ PDF to DOCX Error: {e}")
        raise e

# -----------------------------
# Convert file → PDF
# -----------------------------

def convert_to_pdf(input_file, output_file):
    ext = os.path.splitext(input_file)[1].lower()
    out_dir = os.path.abspath(os.path.dirname(output_file))

    try:
        # --- 1. IMAGE GROUP (Fastest, Internal) ---
        if ext in [".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".gif", ".webp"]:
            img = Image.open(input_file)
            if img.mode != 'RGB':
                img = img.convert('RGB')
            img.save(output_file)
            print(f"✅ Image {ext} converted using Pillow")

        # --- 2. TEXT/WEB GROUP (Internal & Safe) ---
        elif ext in [".txt", ".md"]:
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Arial", size=11)
            with open(input_file, "r", encoding="utf-8", errors="ignore") as f:
                for line in f:
                    # Clean text for PDF compatibility
                    safe_text = line.encode('latin-1', 'replace').decode('latin-1')
                    pdf.cell(0, 10, txt=safe_text, ln=True)
            pdf.output(output_file)
            print(f"✅ {ext} converted using FPDF")

        # --- 3. THE "OFFICE" GROUP (Requires LibreOffice via subprocess) ---
        elif ext in [
            ".docx", ".doc", ".rtf",                       # Word
            ".xls", ".xlsx", ".ods", ".fods",              # Excel
            ".ppt", ".pptx", ".odp",                       # PowerPoint
            ".odt", ".html", ".htm", ".epub",              # Docs/Web
            ".odg", ".otg", ".fodg"                        # Graphics
        ]:
            print(f"🔄 Starting LibreOffice for {ext}...")
            # Headless command to convert
            cmd = [
                LIBRE_OFFICE_PATH,
                "--headless",
                "--convert-to", "pdf",
                input_file,
                "--outdir", out_dir
            ]
            subprocess.run(cmd, capture_output=True, check=True)
            
            # Find and rename the file LibreOffice created
            expected_name = os.path.splitext(os.path.basename(input_file))[0] + ".pdf"
            generated_path = os.path.join(out_dir, expected_name)
            
            if os.path.exists(generated_path):
                if os.path.exists(output_file): os.remove(output_file)
                os.rename(generated_path, output_file)
                print(f"✅ {ext} converted using LibreOffice")
            else:
                raise Exception("LibreOffice finished but no PDF was found.")

    except Exception as e:
        print(f"❌ Error converting {ext}: {e}")
        raise e

# -----------------------------
# Merge PDFs
# -----------------------------
def merge_pdfs(pdf_list, output_file):
    try:
        merger = PdfMerger()
        for pdf in pdf_list:
            merger.append(pdf)
        merger.write(output_file)
        merger.close()
    except Exception as e:
        print(f"❌ Merge Error: {e}")
        raise e

# -----------------------------
# Master function for Flask
# -----------------------------
def master_process(files, mode, target_format):
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    temp_dir = os.path.join(root, "temp_files")
    os.makedirs(temp_dir, exist_ok=True)

    saved_paths = []
    # 1. Save all raw uploads first
    for f in files:
        unique_name = f"raw_{int(time.time())}_{f.filename}"
        path = os.path.join(temp_dir, unique_name)
        f.save(path)
        saved_paths.append(path)

    output_filename = f"result_{int(time.time())}.{target_format}"
    output_path = os.path.join(temp_dir, output_filename)

    # 2. Logic Router
    if mode == "merge":
        # Convert everything to PDF *before* merging
        ready_to_merge = []
        temp_created = []

        for path in saved_paths:
            ext = os.path.splitext(path)[1].lower()
            if ext == ".pdf":
                ready_to_merge.append(path)
            else:
                # Convert non-pdf to temporary pdf
                temp_pdf = path + ".temp.pdf"
                try:
                    convert_to_pdf(path, temp_pdf)
                    ready_to_merge.append(temp_pdf)
                    temp_created.append(temp_pdf)
                except Exception as e:
                    print(f"Skipping {path} due to error: {e}")

        if not ready_to_merge:
            raise Exception("No valid files to merge.")
            
        merge_pdfs(ready_to_merge, output_path)

        # Cleanup intermediate files
        for t in temp_created:
            if os.path.exists(t): os.remove(t)

    else:
        # Single Mode
        input_file = saved_paths[0]
        ext = os.path.splitext(input_file)[1].lower()

        if ext == ".pdf" and target_format != "pdf":
            if target_format == "txt":
                pdf_to_txt(input_file, output_path)
            elif target_format == "docx":
                pdf_to_docx(input_file, output_path)
        else:
            convert_to_pdf(input_file, output_path)

    return output_path


